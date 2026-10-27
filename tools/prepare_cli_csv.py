"""Read authorized CLI XLSX files; create private SaaS CSV candidates without changing source values.
Requires openpyxl for reading only. No network or database writes. Do not publish the output.
Validate candidates with the Backend ProspectCsvParser/preview before any authorized import.
"""
import argparse
import csv
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

FIELDS = {"Name":"name", "Category":"category", "Address":"address", "Phone":"phone", "Email":"email", "Website":"website", "Language":"language", "Status":"status"}

def prepare(sources, output):
    output.mkdir(parents=True, exist_ok=True)
    combined = []
    manifest = []
    columns = None
    ledger = []
    for source in sources:
        digest = hashlib.sha256(source.read_bytes()).hexdigest()
        workbook = load_workbook(source, read_only=True, data_only=False)
        if workbook.sheetnames != ["Businesses"]:
            workbook.close()
            raise ValueError("Unexpected workbook layout; review source mapping explicitly")
        sheet = workbook["Businesses"]
        iterator = sheet.iter_rows()
        headers = [cell.value for cell in next(iterator)]
        if len(set(headers)) != len(headers) or any(h not in FIELDS for h in headers) or "Name" not in headers:
            workbook.close()
            raise ValueError("Unsupported/repeated headers; source mapping requires review")
        mapped = [FIELDS[h] for h in headers]
        if columns is not None and columns != mapped:
            workbook.close()
            raise ValueError("Incompatible source column sets; process sources separately")
        columns = mapped
        records = []
        skipped = 0
        for excel_row, cells in enumerate(iterator, 2):
            if all(cell.value is None for cell in cells):
                skipped += 1
                ledger.append({"source":source.name,"excelRow":excel_row,"candidateRecord":None,"review":["EMPTY_SOURCE_ROW"]})
                continue
            values = []
            warnings = []
            for header, cell in zip(headers, cells):
                if cell.data_type == "f":
                    warnings.append(FIELDS[header]+":FORMULA_REQUIRES_REVIEW")
                elif cell.value is not None and not isinstance(cell.value, str):
                    warnings.append(FIELDS[header]+":NON_TEXT_SOURCE_REQUIRES_REVIEW")
                values.append("" if cell.value is None else str(cell.value))
            records.append(values)
            combined.append(values)
            ledger.append({"source":source.name,"excelRow":excel_row,"candidateRecord":len(combined),"review":warnings})
        workbook.close()
        candidate = output / (source.stem + ".candidate.csv")
        with candidate.open("w",encoding="utf-8",newline="") as stream:
            writer=csv.writer(stream);writer.writerow(columns);writer.writerows(records)
        if hashlib.sha256(source.read_bytes()).hexdigest()!=digest:
            raise ValueError("Source changed during read")
        manifest.append({"source":source.name,"sha256":digest,"bytes":source.stat().st_size,"rows":len(records),"blankRowsRecorded":skipped,"headers":headers,"candidate":candidate.name})
    with (output/"all-records.candidate.csv").open("w",encoding="utf-8",newline="") as stream:
        writer=csv.writer(stream);writer.writerow(columns);writer.writerows(combined)
    (output/"source-manifest.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding="utf-8")
    (output/"row-ledger.json").write_text(json.dumps(ledger,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps({"sources":len(manifest),"records":len(combined),"blankRowsRecorded":sum(m["blankRowsRecorded"] for m in manifest),"cellsRequiringReview":sum(len(r["review"]) for r in ledger if r["candidateRecord"]),"originalsUnchanged":True}))

if __name__ == "__main__":
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument("sources",type=Path,nargs="+")
    parser.add_argument("--output-dir",type=Path,default=Path(".tmp/cli-csv-review"))
    args=parser.parse_args()
    prepare(args.sources,args.output_dir)
